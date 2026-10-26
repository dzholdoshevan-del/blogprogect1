"""Команда Django: собирает ВСЕ тесты проекта, запускает их и записывает
в Word-файл (название, описание, результат, исходный код, ошибки).

Запуск:
    python manage.py export_tests_docx
    python manage.py export_tests_docx --output otchet.docx
    python manage.py export_tests_docx blog          # только тесты приложения blog
    python manage.py export_tests_docx --no-run      # без запуска, только список и код
"""
import inspect
import io
import textwrap
import unittest
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand
from django.test.runner import DiscoverRunner
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = 'Times New Roman'
GREEN, RED, GREY = 'C6EFCE', 'FFC7CE', 'D9D9D9'


# ---------- сбор тестов ----------
def iter_tests(suite):
    for item in suite:
        if isinstance(item, unittest.TestSuite):
            yield from iter_tests(item)
        else:
            yield item


def describe(test):
    """Возвращает (описание, исходный код) теста.
    Описание = docstring или комментарии прямо над def."""
    method = getattr(test, getattr(test, '_testMethodName', ''), None)
    func = getattr(method, '__func__', method)
    if func is None:
        return '', ''
    try:
        lines, start = inspect.getsourcelines(func)
        source = textwrap.dedent(''.join(lines)).rstrip()
        file_lines = Path(inspect.getsourcefile(func)).read_text(encoding='utf-8').splitlines()
    except (OSError, TypeError):
        return inspect.getdoc(func) or '', ''
    comments, i = [], start - 2
    while i >= 0 and file_lines[i].strip().startswith('#'):
        comments.insert(0, file_lines[i].strip().lstrip('#').strip())
        i -= 1
    return inspect.getdoc(func) or ' '.join(comments), source


class Collector(unittest.TextTestResult):
    """Запоминает результат каждого теста."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.records = {}

    def addSuccess(self, test):
        super().addSuccess(test)
        self.records[test.id()] = ('OK', '')

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self.records[test.id()] = ('FAIL', self._exc_info_to_string(err, test))

    def addError(self, test, err):
        super().addError(test, err)
        self.records[test.id()] = ('ERROR', self._exc_info_to_string(err, test))


# ---------- помощники для docx ----------
def set_font(run, size=12, bold=False, italic=False, name=FONT, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(element, fill):
    """Заливка ячейки (tcPr) или абзаца (pPr)."""
    pr = element.get_or_add_tcPr() if hasattr(element, 'get_or_add_tcPr') else element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def para(doc, text='', size=12, bold=False, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), size, bold, italic)
    return p


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    set_font(h.add_run(text), {1: 16, 2: 14, 3: 12}[level], bold=True, color='000000')
    return h


def code_block(doc, text, size=9):
    for line in text.split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        shade(p._p, 'F2F2F2')
        set_font(p.add_run(line or ' '), size, name='Courier New')
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def cell_text(cell, text, bold=False, size=11, fill=None, center=False):
    cell.text = ''
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size, bold)
    if fill:
        shade(cell._tc, fill)


STATUS_RU = {'OK': 'Пройден', 'FAIL': 'Упал', 'ERROR': 'Ошибка', '—': '—'}


class Command(BaseCommand):
    help = 'Запускает все тесты проекта и записывает их в Word-файл.'

    def add_arguments(self, parser):
        parser.add_argument('labels', nargs='*', help='Приложения или тесты (по умолчанию все).')
        parser.add_argument('--output', '-o', default='test_report.docx', help='Имя Word-файла.')
        parser.add_argument('--no-run', action='store_true', help='Не запускать тесты.')

    def handle(self, *args, **options):
        runner = DiscoverRunner(verbosity=0, interactive=False)
        runner.setup_test_environment()
        suite = runner.build_suite(options['labels'] or None)
        tests = list(iter_tests(suite))

        results = {}
        if not options['no_run']:
            old_config = runner.setup_databases()
            try:
                result = unittest.TextTestRunner(
                    stream=io.StringIO(), verbosity=0, resultclass=Collector).run(suite)
                results = result.records
            finally:
                runner.teardown_databases(old_config)
        runner.teardown_test_environment()

        self.build_document(tests, results, options['output'], ran=not options['no_run'])
        self.stdout.write(self.style.SUCCESS(
            f"Готово: {len(tests)} тестов записано в {options['output']}"))

    # ---------- документ ----------
    def build_document(self, tests, results, output, ran):
        doc = Document()
        doc.styles['Normal'].font.name = FONT
        doc.styles['Normal'].font.size = Pt(12)

        para(doc, 'ОТЧЁТ О ТЕСТИРОВАНИИ DJANGO-ПРОЕКТА', 18, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        para(doc, f'Сформирован: {datetime.now():%d.%m.%Y %H:%M}', 12,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

        # группировка по классам, порядок сохраняется
        groups = {}
        for t in tests:
            module = t.__class__.__module__
            groups.setdefault((module, t.__class__.__name__), []).append(t)

        status = lambda t: results.get(t.id(), ('—', ''))[0]
        total = len(tests)
        passed = sum(1 for t in tests if status(t) == 'OK')
        failed = sum(1 for t in tests if status(t) in ('FAIL', 'ERROR'))

        heading(doc, '1. Общая сводка', 1)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_text(table.rows[0].cells[0], 'Показатель', True, fill=GREY)
        cell_text(table.rows[0].cells[1], 'Значение', True, fill=GREY)
        rows = [('Всего тестов', str(total)), ('Тестовых классов', str(len(groups)))]
        if ran:
            rows += [('Пройдено', str(passed)), ('Упало / с ошибкой', str(failed))]
        for k, v in rows:
            r = table.add_row().cells
            cell_text(r[0], k)
            cell_text(r[1], v, center=True)
        para(doc, '', after=6)

        heading(doc, '2. Список тестов по классам', 1)
        for n, ((module, cls), items) in enumerate(groups.items(), 1):
            heading(doc, f'2.{n}. {cls}  ({module})', 2)
            doc_cls = inspect.getdoc(items[0].__class__)
            if doc_cls and not doc_cls.startswith('A class whose instances'):
                para(doc, doc_cls, 12, italic=True)
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            widths = [0.4, 2.6, 2.6, 0.9]
            for i, h in enumerate(['№', 'Тест', 'Что проверяется', 'Результат']):
                cell_text(table.rows[0].cells[i], h, True, fill=GREY)
            for i, t in enumerate(items, 1):
                desc, _ = describe(t)
                st = status(t)
                fill = GREEN if st == 'OK' else RED if st in ('FAIL', 'ERROR') else None
                row = table.add_row().cells
                cell_text(row[0], str(i), center=True)
                cell_text(row[1], getattr(t, '_testMethodName', t.id()), size=10)
                cell_text(row[2], desc or '—', size=10)
                cell_text(row[3], STATUS_RU[st], True, fill=fill, center=True)
            table.autofit = False
            for col, w in zip(table.columns, widths):
                col.width = Pt(w * 72)
            for row in table.rows:
                for c, w in zip(row.cells, widths):
                    c.width = Pt(w * 72)
            para(doc, '', after=6)

        heading(doc, '3. Исходный код тестов', 1)
        for n, ((module, cls), items) in enumerate(groups.items(), 1):
            heading(doc, f'3.{n}. {cls}', 2)
            for t in items:
                _, src = describe(t)
                para(doc, getattr(t, '_testMethodName', t.id()), 12, True, after=2)
                if src:
                    code_block(doc, src)

        errors = [(t, results[t.id()][1]) for t in tests if status(t) in ('FAIL', 'ERROR')]
        if errors:
            heading(doc, '4. Упавшие тесты (подробности)', 1)
            for t, tb in errors:
                para(doc, t.id(), 12, True, after=2)
                code_block(doc, tb.strip())

        doc.save(output)
